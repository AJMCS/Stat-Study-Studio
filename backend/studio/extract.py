"""Automatic text extraction for PDFs, images and plain text, with per-page confidence.

Used for personal uploads and for shared-library files that have no hand transcript yet.

Strategy per PDF page:
  1. Read the text layer. LaTeX PDFs often carry math glyphs that map to the wrong characters
     (theta comes out as a check mark, pi as an arrow), so the text layer is checked for those.
  2. If the page has no usable text, or the text layer looks corrupted, render the page and OCR it.
  3. Report a confidence and a human-readable warning for pages that may be missing symbols.
"""

from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import numpy as np

# Characters that LaTeX math fonts commonly decode to when the PDF has no proper mapping.
_SUSPECT = re.compile("[←-⇿✀-➿-■-◿�]")
_ALPHA = re.compile(r"[A-Za-z]")
_MATHY = re.compile(r"[=+\-<>/^_]")

MIN_TEXT_CHARS = 40
OCR_DPI = 150
MAX_PDF_PAGES = 400
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
ACCEPTED_SUFFIXES = {".pdf", ".docx", ".doc", ".txt"} | IMAGE_SUFFIXES

Progress = Callable[[int, int], None]


class ParseError(Exception):
    """The file cannot be used. The message is written for the learner and says what to do next."""

_ocr_engine = None


def ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr import RapidOCR

        _ocr_engine = RapidOCR()
    return _ocr_engine


@dataclass
class PageText:
    page: int
    text: str
    method: str  # text-layer | ocr | text-file
    confidence: float  # 0..1
    flags: list[str] = field(default_factory=list)

    @property
    def warning(self) -> str | None:
        if "text-layer-corrupted" in self.flags:
            return (
                "Math on this page was read automatically. Fractions, sums and subscripts may be missing or out of "
                "order. Check it against the original page."
            )
        if "low-ocr-confidence" in self.flags:
            return "This page was read with low confidence. The text may be incomplete or inaccurate."
        if "little-text" in self.flags:
            return "Very little text was found on this page."
        return None


def normalize(text: str) -> str:
    """NFKC folds math-italic letters (U+1D44E and friends) back to plain letters and Greek."""
    return unicodedata.normalize("NFKC", text)


def suspect_ratio(text: str) -> float:
    alpha = len(_ALPHA.findall(text))
    if alpha == 0:
        return 1.0 if _SUSPECT.search(text) else 0.0
    return len(_SUSPECT.findall(text)) / alpha


def _ocr_image(arr: np.ndarray) -> tuple[str, float]:
    res = ocr_engine()(arr)
    txts = list(res.txts or [])
    scores = list(res.scores or [])
    return "\n".join(txts), (float(np.mean(scores)) if scores else 0.0)


def extract_pdf(path: Path, force_ocr: bool = False, progress: Progress | None = None) -> list[PageText]:
    import pymupdf
    from PIL import Image

    pages: list[PageText] = []
    try:
        doc = pymupdf.open(path)
    except Exception as e:
        raise ParseError("This PDF could not be opened. It may be damaged. Try opening it and saving a fresh copy, then upload that.") from e
    if doc.needs_pass:
        raise ParseError("This PDF is password-protected. Remove the password (open it, then print or save a copy without one) and upload that.")
    if len(doc) == 0:
        raise ParseError("This PDF has no pages.")
    if len(doc) > MAX_PDF_PAGES:
        raise ParseError(f"This PDF has {len(doc)} pages; the limit is {MAX_PDF_PAGES}. Split it into smaller files and upload those.")
    for i, page in enumerate(doc, 1):
        if progress:
            progress(i - 1, len(doc))
        layer = normalize(page.get_text()).strip()
        ratio = suspect_ratio(layer)
        usable = len(layer) >= MIN_TEXT_CHARS and ratio < 0.004
        if usable and not force_ocr:
            pages.append(PageText(i, layer, "text-layer", 0.95))
            continue
        pix = page.get_pixmap(dpi=OCR_DPI)
        img = np.array(Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB"))
        text, conf = _ocr_image(img)
        text = normalize(text)
        flags = []
        if len(layer) >= MIN_TEXT_CHARS and ratio >= 0.004:
            flags.append("text-layer-corrupted")
            conf = min(conf, 0.85)  # OCR reads the letters, not the math layout
        if len(text) < MIN_TEXT_CHARS:
            flags.append("little-text")
        if conf < 0.8:
            flags.append("low-ocr-confidence")
        pages.append(PageText(i, text, "ocr", conf, flags))
    if progress:
        progress(len(doc), len(doc))
    return pages


def extract_image(path: Path) -> list[PageText]:
    from PIL import Image

    try:
        img = np.array(Image.open(path).convert("RGB"))
    except Exception as e:
        raise ParseError("This image could not be opened. It may be damaged or not a supported picture. Try saving it as a PNG or JPG.") from e
    text, conf = _ocr_image(img)
    text = normalize(text)
    flags = []
    if len(text) < MIN_TEXT_CHARS:
        flags.append("little-text")
    if conf < 0.8:
        flags.append("low-ocr-confidence")
    return [PageText(1, text, "ocr", conf, flags)]


def extract_text_file(path: Path) -> list[PageText]:
    data = path.read_bytes()
    if b"\x00" in data[:4096]:
        raise ParseError("This file does not look like plain text (it contains binary data). Upload a PDF, Word file, or a real .txt file.")
    raw = normalize(data.decode("utf-8", errors="replace")).strip()
    paras = [p.strip() for p in re.split(r"\n{2,}", raw) if p.strip()]
    # Group paragraphs into pseudo-pages so citations can point somewhere.
    pages, cur, n = [], [], 1
    for p in paras:
        cur.append(p)
        if sum(len(x) for x in cur) > 2500:
            pages.append(PageText(n, "\n\n".join(cur), "text-file", 1.0))
            cur, n = [], n + 1
    if cur:
        pages.append(PageText(n, "\n\n".join(cur), "text-file", 1.0))
    return pages


def _pseudo_pages(blocks: list[str]) -> list[PageText]:
    """Word and text files have no fixed pages; group them so a citation can still point somewhere."""
    pages, cur, n = [], [], 1
    for b in blocks:
        cur.append(b)
        if sum(len(x) for x in cur) > 2500:
            pages.append(PageText(n, "\n\n".join(cur), "text-file", 1.0))
            cur, n = [], n + 1
    if cur:
        pages.append(PageText(n, "\n\n".join(cur), "text-file", 1.0))
    return pages


def extract_docx(path: Path) -> list[PageText]:
    try:
        import docx

        document = docx.Document(str(path))
    except Exception as e:
        raise ParseError("This Word file could not be opened. It may be damaged, or not a real .docx. Try saving it again from Word, or as a PDF.") from e
    blocks: list[str] = []
    for para in document.paragraphs:
        text = normalize(para.text).strip()
        if not text:
            continue
        style = (para.style.name or "") if para.style is not None else ""
        blocks.append(f"# {text}" if style.lower().startswith("heading") or style.lower() == "title" else text)
    for table in document.tables:
        rows = [" | ".join(normalize(c.text).strip() for c in row.cells) for row in table.rows]
        if any(r.strip(" |") for r in rows):
            blocks.append("\n".join(rows))
    return _pseudo_pages(blocks)


def extract_doc(path: Path) -> list[PageText]:
    """Legacy .doc has no pure-Python reader. Use LibreOffice if the learner happens to have it; otherwise explain."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        with tempfile.TemporaryDirectory() as out:
            try:
                subprocess.run([soffice, "--headless", "--convert-to", "docx", "--outdir", out, str(path)], check=True, capture_output=True, timeout=120)
                converted = next(Path(out).glob("*.docx"), None)
                if converted:
                    return extract_docx(converted)
            except (subprocess.SubprocessError, OSError):
                pass
    raise ParseError(
        "Older Word files (.doc) can't be read here. In Word, choose File > Save As and pick Word Document (.docx) or PDF, then upload that copy."
    )


def extract(path: Path, force_ocr: bool = False, progress: Progress | None = None) -> list[PageText]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path, force_ocr, progress)
    if suffix == ".txt":
        return extract_text_file(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".doc":
        return extract_doc(path)
    if suffix in IMAGE_SUFFIXES:
        return extract_image(path)
    raise ParseError(
        f'"{suffix or "this file"}" is not a supported type. Upload a PDF, Word file (.docx), .txt, or an image (PNG, JPG, TIFF, BMP, WebP).'
    )


def to_transcript(pages: list[PageText], source: str, kind: str) -> str:
    """Render extracted pages in the same page-marked Markdown the hand transcripts use."""
    confs = [p.confidence for p in pages]
    head = [
        "---",
        f"source: {source}",
        "method: automatic extraction (text layer, with OCR where the layer was missing or corrupted); not hand-verified",
        f"pages: {len(pages)}",
        f"kind: {kind}",
        f"confidence: {np.mean(confs):.2f}" if confs else "confidence: 0",
        "---",
        "",
    ]
    body = []
    for p in pages:
        flags = f" flags={','.join(p.flags)}" if p.flags else ""
        body.append(f"<!-- page {p.page} -->")
        body.append(f"<!-- extracted via {p.method} confidence={p.confidence:.2f}{flags} -->")
        body.append(p.text.strip())
        body.append("")
    return "\n".join(head + body)
