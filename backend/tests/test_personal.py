import io
import time
from pathlib import Path

import pymupdf
import pytest
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont

from studio.app import create_app

PORT = 8765
ORIGIN = {"Origin": f"http://127.0.0.1:{PORT}"}

LECTURE = "The posterior distribution describes what we know about a parameter after observing the data."
MLE_TIP = "Always check that the second derivative of the log likelihood is negative at the candidate maximum."


# ---------------------------------------------------------------- fixtures and builders
@pytest.fixture()
def client(tmp_path):
    app = create_app(data_dir=tmp_path / "data", static_dir=tmp_path / "x", port=PORT)
    with TestClient(app, base_url=f"http://127.0.0.1:{PORT}") as c:
        c.data_dir = tmp_path / "data"
        c.app = app
        yield c


def text_pdf(pages: list[str]) -> bytes:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(60, 60, 540, 780), text, fontsize=12)
    return doc.tobytes()


def picture_of(text: str, size=(1400, 500)) -> Image.Image:
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default(size=44)
    y = 40
    for line in text.split("\n"):
        draw.text((40, y), line, fill="black", font=font)
        y += 80
    return img


def png_bytes(text: str) -> bytes:
    buf = io.BytesIO()
    picture_of(text).save(buf, "PNG")
    return buf.getvalue()


def scanned_pdf(text: str) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(40, 60, 560, 260), stream=png_bytes(text))
    return doc.tobytes()


def docx_bytes() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Study tips", level=1)
    d.add_paragraph(MLE_TIP)
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "Concept"
    t.rows[0].cells[1].text = "Reminder"
    t.rows[1].cells[0].text = "Conjugate prior"
    t.rows[1].cells[1].text = "Posterior stays in the same family as the prior"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def upload(client, name: str, data: bytes, collection="general", replace=None, expect=202):
    form = {"collectionId": collection}
    if replace:
        form["replaceId"] = replace
    r = client.post("/api/uploads", files={"file": (name, data)}, data=form, headers=ORIGIN)
    assert r.status_code == expect, r.text
    return r


def finish(client, upload_id: str, timeout=90) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = client.get(f"/api/uploads/{upload_id}").json()
        if job["status"] in {"ready", "failed"}:
            return job
        time.sleep(0.1)
    raise AssertionError("upload did not finish")


def add(client, name, data, **kw):
    job = finish(client, upload(client, name, data, **kw).json()["uploadId"])
    assert job["status"] == "ready", job
    return job


def personal(client, q, **params):
    r = client.get("/api/search", params={"q": q, **params})
    assert r.status_code == 200, r.text
    return r.json()["personal"]


def put(client, name, value):
    assert client.put(f"/api/state/{name}", json=value, headers=ORIGIN).status_code == 200


def collection(client, cid="hw3", name="Homework 3"):
    cols = client.get("/api/state").json()["state"]["collections"]
    put(client, "collections", cols + [{"id": cid, "name": name, "consent": False, "searchByDefault": False}])


# ---------------------------------------------------------------- the main formats
def test_text_file_becomes_searchable_with_a_full_citation(client):
    job = add(client, "my notes.txt", f"{MLE_TIP}\n\nA second paragraph about priors.".encode())
    f = job["file"]
    assert f["filename"] == "my notes.txt" and f["version"] == 1 and f["collectionId"] == "general" and f["removed"] is False
    hit = personal(client, "second derivative of the log likelihood")[0]
    assert hit["scope"] == "personal" and hit["collection"] == "General" and hit["filename"] == "my notes.txt" and hit["page"] == 1
    assert hit["version"] == 1 and hit["fileId"] == f["id"] and hit["ocr"] is False
    page = client.get(f"/api/files/{f['id']}/pages/1").json()
    assert hit["highlight"] and hit["highlight"] in page["text"]


def test_text_pdf_keeps_page_numbers_and_serves_page_pictures_and_original(client):
    data = text_pdf(["Page one is about priors and likelihoods in general.", "Page two explains the standard error of an estimator."])
    f = add(client, "lecture.pdf", data)["file"]
    assert f["pages"] == 2 and f["ocr"] is False
    hit = personal(client, "standard error of an estimator")[0]
    assert hit["page"] == 2
    img = client.get(f"/api/files/{f['id']}/pages/2/image")
    assert img.status_code == 200 and img.headers["content-type"] == "image/png" and img.content[:4] == b"\x89PNG"
    original = client.get(f"/api/files/{f['id']}/original")
    assert original.content == data and original.headers["content-type"] == "application/pdf"
    assert "inline" in original.headers["content-disposition"] and original.headers["x-content-type-options"] == "nosniff"


def test_docx_text_and_tables_are_indexed_and_have_no_page_picture(client):
    f = add(client, "tips.docx", docx_bytes())["file"]
    assert personal(client, "second derivative negative candidate maximum")
    assert personal(client, "conjugate prior posterior stays in the same family")
    assert client.get(f"/api/files/{f['id']}/pages/1/image").status_code == 404


def test_scanned_pdf_and_photo_are_read_with_ocr_and_marked_as_ocr(client):
    scan = add(client, "scan.pdf", scanned_pdf("The posterior mean is the Bayes estimate\nunder squared error loss"))["file"]
    photo = add(client, "photo.png", png_bytes("Maximum likelihood estimation chooses\nthe parameter that makes the data likely"))["file"]
    assert scan["ocr"] is True and photo["ocr"] is True
    hit = personal(client, "posterior mean Bayes estimate squared error loss")[0]
    assert hit["fileId"] == scan["id"] and hit["ocr"] is True
    hit = personal(client, "maximum likelihood estimation chooses the parameter")[0]
    assert hit["fileId"] == photo["id"] and hit["ocr"] is True
    page = client.get(f"/api/files/{photo['id']}/pages/1").json()
    assert page["method"] == "ocr" and page["hasImage"] is True
    assert client.get(f"/api/files/{photo['id']}/pages/1/image").content[:4] == b"\x89PNG"


def test_upload_report_flags_pages_that_need_a_second_look(client):
    job = add(client, "scan.pdf", scanned_pdf("The posterior mean is the Bayes estimate\nunder squared error loss"))
    r = job["report"]
    assert r["pages"] == 1 and r["ocrPages"] == 1 and isinstance(r["flaggedPages"], list)


# ---------------------------------------------------------------- rejection: never added, never searchable
def assert_nothing_added(client):
    assert client.get("/api/files").json() == []
    assert personal(client, "posterior distribution parameter") == []


@pytest.mark.parametrize(
    "name,data,message",
    [
        ("program.exe", b"MZ....", "not a supported type"),
        ("empty.txt", b"", "empty"),
        ("legacy.doc", b"\xd0\xcf\x11\xe0 old word file", "Save As"),
    ],
)
def test_rejected_immediately_with_a_reason(client, name, data, message):
    r = upload(client, name, data, expect=422)
    assert message in r.json()["detail"]
    assert_nothing_added(client)


@pytest.mark.parametrize(
    "name,data,message",
    [
        ("broken.pdf", b"this is not a pdf at all", "could not be opened"),
        ("binary.txt", bytes(range(256)) * 20, "binary"),
        ("blank.png", None, "No readable text"),
        ("fake.docx", b"PK not really a docx", "could not be opened"),
        ("bad.png", b"\x89PNG not an image", "could not be opened"),
    ],
)
def test_unreadable_files_fail_with_a_reason_and_are_not_added(client, name, data, message):
    if name == "blank.png":
        buf = io.BytesIO()
        Image.new("RGB", (800, 600), "white").save(buf, "PNG")
        data = buf.getvalue()
    job = finish(client, upload(client, name, data).json()["uploadId"])
    assert job["status"] == "failed" and message in job["error"], job
    assert job["file"] is None
    assert_nothing_added(client)


def test_password_protected_pdf_is_rejected_with_instructions(client):
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret notes about the posterior distribution")
    locked = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="owner")
    job = finish(client, upload(client, "locked.pdf", locked).json()["uploadId"])
    assert job["status"] == "failed" and "password" in job["error"]
    assert_nothing_added(client)


def test_unknown_collection_is_rejected(client):
    r = upload(client, "n.txt", b"hello there friends " * 5, collection="nope", expect=404)
    assert "no longer exists" in r.json()["detail"]


def test_oversize_upload_is_refused_without_keeping_anything(client, monkeypatch):
    monkeypatch.setattr("studio.app.MAX_BYTES", 1024)
    r = upload(client, "big.txt", b"x" * 5000, expect=413)
    assert "larger than" in r.json()["detail"]
    assert list((client.data_dir / "personal" / "_incoming").iterdir()) == []


# ---------------------------------------------------------------- duplicates and versions
def test_duplicates_are_refused_with_guidance(client):
    add(client, "notes.txt", MLE_TIP.encode())
    same = upload(client, "copy.txt", MLE_TIP.encode(), expect=422).json()["detail"]
    assert "already in" in same
    renamed = upload(client, "NOTES.txt", b"A different file that happens to share a name, about priors.", expect=422).json()["detail"]
    assert "Replace" in renamed
    assert len(client.get("/api/files").json()) == 1


def test_replacing_creates_a_new_version_and_only_the_latest_is_searched(client):
    v1 = add(client, "notes.txt", b"Old advice: the posterior is always a normal distribution whatever the prior.")["file"]
    v2 = add(client, "notes.txt", b"New advice: a conjugate prior keeps the posterior in the same family as the prior.", replace=v1["id"])["file"]
    assert v2["version"] == 2 and v2["supersedes"] == v1["id"]
    assert personal(client, "always a normal distribution whatever the prior") == [] or all(h["fileId"] == v2["id"] for h in personal(client, "always a normal distribution whatever the prior"))
    assert personal(client, "conjugate prior keeps the posterior in the same family")[0]["fileId"] == v2["id"]
    old_page = client.get(f"/api/files/{v1['id']}/pages/1")  # still resolvable for past activity
    assert old_page.status_code == 200 and "Old advice" in old_page.json()["text"] and old_page.json()["version"] == 1
    assert upload(client, "notes.txt", b"New advice: a conjugate prior keeps the posterior in the same family as the prior.", replace=v2["id"], expect=422)


def test_replace_needs_a_current_file(client):
    f = add(client, "a.txt", b"Some study notes about the likelihood function and data.")["file"]
    add(client, "a.txt", b"Updated study notes about the likelihood function and data, version two.", replace=f["id"])
    assert upload(client, "a.txt", b"third version of the notes about likelihood functions here", replace=f["id"], expect=404)


# ---------------------------------------------------------------- removal
def test_removed_file_leaves_search_and_disk_but_keeps_a_record(client):
    f = add(client, "notes.pdf", text_pdf([LECTURE]))["file"]
    assert personal(client, "what we know about a parameter after observing the data")
    assert client.delete(f"/api/files/{f['id']}", headers=ORIGIN).status_code == 200
    assert personal(client, "what we know about a parameter after observing the data") == []
    assert client.get(f"/api/files/{f['id']}/pages/1").status_code == 404
    assert client.get(f"/api/files/{f['id']}/original").status_code == 404
    record = client.get("/api/files").json()[0]
    assert record["removed"] is True and record["filename"] == "notes.pdf"
    assert not (client.data_dir / "personal" / f["id"]).exists()
    assert client.delete(f"/api/files/{f['id']}", headers=ORIGIN).status_code == 404


def test_removing_the_latest_version_removes_older_versions_too(client):
    v1 = add(client, "n.txt", b"Version one talks about priors and the posterior distribution in detail.")["file"]
    v2 = add(client, "n.txt", b"Version two talks about likelihood and the maximum likelihood estimator.", replace=v1["id"])["file"]
    client.delete(f"/api/files/{v2['id']}", headers=ORIGIN)
    assert all(f["removed"] for f in client.get("/api/files").json())
    assert not (client.data_dir / "personal" / v1["id"]).exists()
    assert client.get(f"/api/files/{v1['id']}/pages/1").status_code == 404


# ---------------------------------------------------------------- collections
def test_general_is_always_searched_and_named_collections_only_when_chosen(client):
    collection(client)
    add(client, "g.txt", b"General notes about the standard error and the sampling distribution of the mean.")
    add(client, "h.txt", b"Homework notes about the Bayes estimator under absolute error loss and the posterior median.", collection="hw3")
    q = "Bayes estimator absolute error posterior median"
    assert all(h["collectionId"] != "hw3" for h in personal(client, q))  # a named collection is not searched unless chosen
    named = personal(client, q, collection="hw3")
    assert named and named[0]["collection"] == "Homework 3" and named[0]["collectionId"] == "hw3"
    assert personal(client, "standard error sampling distribution of the mean", collection="hw3")  # General still included


# ---------------------------------------------------------------- moved / missing originals
def test_missing_original_excludes_the_file_until_reconnected(client):
    data = MLE_TIP.encode()
    f = add(client, "notes.txt", data)["file"]
    (client.data_dir / "personal" / f["id"] / "original.txt").unlink()
    assert personal(client, "second derivative log likelihood negative") == []
    assert client.get("/api/files").json()[0]["linkBroken"] is True
    assert client.get(f"/api/files/{f['id']}/original").status_code == 404
    wrong = client.post(f"/api/files/{f['id']}/reconnect", files={"file": ("x.txt", b"different content entirely")}, headers=ORIGIN)
    assert wrong.status_code == 422 and "isn't the same file" in wrong.json()["detail"]
    ok = client.post(f"/api/files/{f['id']}/reconnect", files={"file": ("notes.txt", data)}, headers=ORIGIN)
    assert ok.status_code == 200
    assert personal(client, "second derivative log likelihood negative")
    assert "linkBroken" not in client.get("/api/files").json()[0]


# ---------------------------------------------------------------- model availability
def test_works_without_the_embedding_model_and_backfills_later(client, monkeypatch):
    embedder = client.app.state.embedder
    real_ensure = embedder.ensure
    monkeypatch.setattr(embedder, "ensure", lambda: None)
    f = add(client, "n.txt", b"The posterior mean minimises squared error loss, so it is the Bayes estimate.")["file"]
    rows = client.app.state.db.read("SELECT embedding FROM personal_chunks WHERE file_id=?", (f["id"],))
    assert rows and all(r["embedding"] is None for r in rows)
    assert personal(client, "Bayes estimate squared error loss")  # keyword search still works
    assert client.get("/api/search", params={"q": "Bayes estimate"}).json()["mode"] == "keyword"
    monkeypatch.setattr(embedder, "ensure", real_ensure)
    assert real_ensure() is not None  # model is on this machine
    assert personal(client, "Bayes estimate squared error loss")
    rows = client.app.state.db.read("SELECT embedding FROM personal_chunks WHERE file_id=?", (f["id"],))
    assert all(r["embedding"] is not None for r in rows)


# ---------------------------------------------------------------- housekeeping and safety
def test_nothing_personal_is_ever_returned_as_a_shared_result(client):
    add(client, "n.txt", b"Personal remark about conjugate priors and the beta distribution.")
    shared = client.get("/api/search", params={"q": "conjugate priors beta distribution"}).json()["results"]
    assert shared and all(h["scope"] == "shared" for h in shared)


def test_delete_all_data_erases_uploaded_files_too(client):
    f = add(client, "n.txt", b"Remember that the likelihood is not a probability distribution over theta.")["file"]
    assert (client.data_dir / "personal" / f["id"]).exists()
    assert client.delete("/api/data", headers=ORIGIN).status_code == 200
    assert client.get("/api/files").json() == []
    assert not (client.data_dir / "personal" / f["id"]).exists()
    assert personal(client, "likelihood probability distribution over theta") == []


def test_folders_without_a_record_are_cleaned_up_on_start(tmp_path):
    kwargs = dict(data_dir=tmp_path / "d", static_dir=tmp_path / "x", port=PORT)
    with TestClient(create_app(**kwargs), base_url=f"http://127.0.0.1:{PORT}") as c:
        stray = tmp_path / "d" / "personal" / "abandoned"
        stray.mkdir(parents=True)
        (stray / "original.pdf").write_bytes(b"x")
    with TestClient(create_app(**kwargs), base_url=f"http://127.0.0.1:{PORT}"):
        assert not stray.exists()


def test_file_ids_cannot_be_used_to_reach_other_files(client):
    for bad in ["..", "../..", "%2e%2e", "nope", "../studio.sqlite3"]:
        assert client.get(f"/api/files/{bad}/original").status_code in (404, 405)
        assert client.get(f"/api/files/{bad}/pages/1").status_code in (404, 405)


def test_uploads_and_removals_require_the_app_origin(client):
    assert client.post("/api/uploads", files={"file": ("a.txt", b"hello hello hello hello")}, data={"collectionId": "general"}).status_code == 403
    assert client.delete("/api/files/abc").status_code == 403


def test_the_files_slice_is_owned_by_the_service(client):
    assert client.put("/api/state/files", json=[], headers=ORIGIN).status_code == 404
    assert client.post("/api/import", json={"files": []}, headers=ORIGIN).status_code == 422


def test_unknown_upload_id_is_a_404(client):
    assert client.get("/api/uploads/nope").status_code == 404


def test_database_upgrades_from_the_first_schema(tmp_path):
    import sqlite3

    from studio.db import SCHEMA_V1, Database, _statements

    path = tmp_path / "old.sqlite3"
    con = sqlite3.connect(path)
    for stmt in _statements(SCHEMA_V1):
        con.execute(stmt)
    con.execute("INSERT INTO collections (position, id, name, consent, search_by_default) VALUES (0,'general','General',0,1)")
    con.execute(
        "INSERT INTO files (position,id,collection_id,filename,version,uploaded_at,pages,ocr,removed) VALUES (0,'f1','general','old.pdf',1,'x',3,0,0)"
    )
    con.execute("PRAGMA user_version=1")
    con.commit()
    con.close()
    db = Database(path)
    state = db.get_state()
    assert state["files"][0]["filename"] == "old.pdf" and state["files"][0]["flaggedPages"] == 0
    assert db.read("PRAGMA user_version")[0][0] == 4
    db.read("SELECT ai_scored FROM attempts")
    db.read("SELECT * FROM personal_chunks")
    db.close()
